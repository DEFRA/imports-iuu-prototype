#!/usr/bin/env python3

import argparse
import io
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Pt
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas


VERSION_ROOT = Path(__file__).resolve().parents[1]
TARGET = VERSION_ROOT / 'data' / 'sample-documents'
TEMPLATES = TARGET / 'templates'
GENERATED_FILENAMES = {
    'CL-2026-44-000079-N.pdf',
    'CATCH.PS.PT.2026.0001149 (Exp. 0125-26-GB).pdf',
    'IUU packing list.pdf'
}

IMPORTER = 'New England Seafood International Ltd, Genesis Way, Healing, Grimsby DN37 9TU, United Kingdom'

CATCH_CERTIFICATES = [
    {
        'filename': 'FRA-2026-CSP-000205.pdf',
        'number': 'FRA 2026 CSP 000205',
        'authority': 'Centre National de Surveillance des Peches, 40 Avenue Louis Bougo, BP 48, 56410 Etel, France',
        'vessel': 'PENDRUC', 'flag': 'FRANCE - CONCARNEAU - CC932207', 'call_sign': 'FIXF', 'imo': '9741102',
        'licence': 'CTOI-1302/000205', 'gear': 'PS (01.1 Purse Seine)',
        'species': 'Skipjack tuna (Katsuwonus pelamis) - 118,000 kg; Yellowfin tuna (Thunnus albacares) - 34,000 kg', 'code': '030343; 030342',
        'area_dates': 'FAO 51 - Indian Ocean Western; 1-18 July 2026',
        'weights': ('152,000', '152,000', '152,000'),
        'exporter': 'Compagnie Francaise du Thon Oceanique, 11 Rue des Sardiniers, 29900 Concarneau, France',
        'transport': ('France - Concarneau', 'BL-FRA-2026-0205', 'MSCU2205101')
    },
    {
        'filename': 'CATCH.CC.FR.2026.0000148 for FRA.2025.CSP.000518.pdf',
        'number': 'FRA 2026 CSP 100124',
        'authority': 'Centre National de Surveillance des Peches, 40 Avenue Louis Bougo, BP 48, 56410 Etel, France',
        'vessel': 'BERNICA', 'flag': 'FRA - DZAOUDZI - DI 929727', 'call_sign': 'FLTZ', 'imo': '9600853',
        'licence': 'CTOI-1302/000201', 'gear': 'PS (01.1 Purse Seine)',
        'species': 'Skipjack tuna (Katsuwonus pelamis)', 'code': '030343',
        'area_dates': 'FAO 51 - Indian Ocean Western; 14-21 June 2026',
        'weights': ('175,564', '174,564', '173,564'),
        'exporter': 'SAPMER S.A., Darse de Peche, 97420 Le Port, Reunion, France',
        'transport': ('France (Reunion) - Le Port', 'BL-FRA-2026-1457', 'MSCU1000333')
    },
    {
        'filename': 'ESP.SGCI.AI.2025.944.pdf',
        'number': 'ESP/SGCI/AI/2026/101',
        'authority': 'Secretaria General de Pesca, C/ Velazquez 147, 28002 Madrid, Spain',
        'vessel': 'ELAI ALAI', 'flag': 'SPAIN - BERMEO - 3BI-2-1-93', 'call_sign': 'EAIW', 'imo': '9046966',
        'licence': 'P0099-6/2026', 'gear': 'PS (01.1 Purse Seine)',
        'species': 'Skipjack tuna (Katsuwonus pelamis)', 'code': '030343',
        'area_dates': 'FAO 51 - Indian Ocean; 1-18 June 2026',
        'weights': ('120,000', '119,000', '118,000'),
        'exporter': 'ECHEBASTAR FLEET SLU, Muelle Erroxape S/N, 48370 Bermeo, Spain',
        'transport': ('Seychelles - Port Victoria', 'BL-ESP-2026-0001', 'MSCU1000030')
    },
    {
        'filename': 'FRA-2025-CSP-000472.pdf',
        'number': 'FRA 2026 CSP 000101',
        'authority': 'Centre National de Surveillance des Peches, 40 Avenue Louis Bougo, BP 48, 56410 Etel, France',
        'vessel': 'BERNICA', 'flag': 'FRA - DZAOUDZI - DI 929727', 'call_sign': 'FLTZ', 'imo': '9600853',
        'licence': 'CTOI-1302/000201', 'gear': 'PS (01.1 Purse Seine)',
        'species': 'Skipjack tuna (Katsuwonus pelamis)', 'code': '030343',
        'area_dates': 'FAO 51 - Indian Ocean Western; 4-21 June 2026',
        'weights': ('80,500', '79,200', '78,000'),
        'exporter': 'SAPMER S.A., Darse de Peche, 97420 Le Port, Reunion, France',
        'transport': ('France (Reunion) - Le Port', 'BL-FRA-2026-0002', 'MSCU1000030')
    },
    {
        'filename': 'SYC-SFA-10-2025-SW0454.pdf',
        'number': 'SYC/SFA/10/2026-SW0454',
        'authority': 'Seychelles Fishing Authority, Fishing Port, Victoria, Mahe, Seychelles',
        'vessel': 'OCEAN VOYAGER', 'flag': 'SEYCHELLES - VICTORIA - SYC7721', 'call_sign': 'OVGR', 'imo': '9800006',
        'licence': 'SYC-171702', 'gear': 'PS (01.1 Purse Seine)',
        'species': 'Skipjack tuna (Katsuwonus pelamis)', 'code': '030343',
        'area_dates': 'FAO 51 - Indian Ocean; 12-28 June 2026',
        'weights': ('66,200', '65,000', '64,000'),
        'exporter': 'Indian Ocean Tuna Exports Ltd, New Port, Victoria, Mahe, Seychelles',
        'transport': ('Seychelles - Port Victoria', 'BL-SYC-2026-0003', 'MSCU1000030')
    },
    {
        'filename': 'CL-2026-44-000079-N.pdf',
        'number': 'CL-2026-44-000079-N',
        'authority': 'Servicio Nacional de Pesca y Acuicultura, Victoria 2832, Valparaiso, Chile',
        'vessel': 'PACIFIC DAWN', 'flag': 'CHILE - VALPARAISO - CL55092', 'call_sign': 'PDWN', 'imo': '9800005',
        'licence': 'CL-PS-2026-079', 'gear': 'PS (01.1 Purse Seine)',
        'species': 'Skipjack tuna (Katsuwonus pelamis)', 'code': '030343',
        'area_dates': 'FAO 87 - South East Pacific; 20 June-8 July 2026',
        'weights': ('102,000', '100,800', '100,000'),
        'exporter': 'Pacific Seafood Chile S.A., Muelle Prat 887, Valparaiso, Chile',
        'transport': ('Chile - Valparaiso', 'BL-CL-2026-0004', 'MSCU1000030')
    }
]

PROCESSING_STATEMENTS = [
    {
        'filename': 'CATCH.PS.PT.2026.0001021 - original processing statement.pdf',
        'number': 'CATCH.PS.PT.2026.0001021',
        'description': 'Frozen skipjack tuna loins - CN 1604 14 26',
        'certificates': [('FRA 2026 CSP 100124', 'BERNICA - France', '31 August 2026', 'Skipjack tuna', '173,564', '165,000', '148,500')]
    },
    {
        'filename': 'CATCH.PS.PT.2026.0001149 (Exp. 0125-26-GB).pdf',
        'number': 'CATCH.PS.PT.2026.0001149',
        'description': 'Frozen skipjack tuna loins - CN 1604 14 26',
        'certificates': [
            ('ESP/SGCI/AI/2026/101', 'ELAI ALAI - Spain', '9 December 2026', 'Skipjack tuna', '118,000', '114,000', '102,600'),
            ('FRA 2026 CSP 000101', 'BERNICA - France', '9 December 2026', 'Skipjack tuna', '78,000', '74,000', '66,600'),
            ('CL-2026-44-000079-N', 'PACIFIC DAWN - Chile', '9 December 2026', 'Skipjack tuna', '100,000', '98,000', '88,200')
        ]
    }
]

BILLS_OF_LADING = [
    {
        'filename': 'IUU packing list.pdf',
        'number': 'BOL-2026-55190',
        'shipper': 'Dakar Ocean Exports SA, Port de Dakar, Senegal',
        'consignee': 'Atlantic Seafoods Ltd, Felixstowe, United Kingdom',
        'container': 'MSCU 7391842',
        'cargo': 'Frozen yellowfin tuna',
        'weight': '24,800 kg',
        'departure': 'Port of Dakar, Senegal',
        'destination': 'Port of Felixstowe, United Kingdom',
        'issued': '16 July 2026'
    }
]


def append_value(cell, value):
    paragraph = cell.add_paragraph()
    run = paragraph.add_run(value)
    run.bold = True
    run.font.size = Pt(7)


def append_to_matching_cell(table, label, value):
    seen = set()
    for row in table.rows:
        for cell in row.cells:
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            if label in cell.text:
                append_value(cell, value)
                return
    raise ValueError(f'Could not find template cell: {label}')


def unique_cells(row):
    cells = []
    seen = set()
    for cell in row.cells:
        if cell._tc not in seen:
            seen.add(cell._tc)
            cells.append(cell)
    return cells


def build_catch_certificate(data, output):
    document = Document(TEMPLATES / 'Catch Cert Sample.docx')
    main, validation, importer, _, transport = document.tables
    values = {
        'Document number': data['number'],
        'Validating authority': data['authority'],
        '1. Name': data['authority'],
        '2. Fishing vessel name': data['vessel'],
        'Flag - home port and registration number': data['flag'],
        'Call sign': data['call_sign'],
        'IMO number': data['imo'],
        'Fishing licence no': data['licence'],
        'Fishing gear': data['gear'],
        '5. Name of master': f"Master of {data['vessel']} - signed electronically for sample",
    }
    for label, value in values.items():
        append_to_matching_cell(main, label, value)
    species_cells = unique_cells(main.rows[8])
    for cell, value in zip(species_cells, [data['species'], data['code'], data['area_dates'], *data['weights']]):
        append_value(cell, value)
    append_to_matching_cell(validation, '8. Name and address of exporter', data['exporter'] + '\nSigned electronically - 10 December 2026')
    append_to_matching_cell(validation, '9. Flag State authority validation', data['authority'] + '\nValidated electronically - 10 December 2026')
    append_to_matching_cell(importer, 'Company, name, address, EORI', IMPORTER + '\nEORI GB987654321000')
    append_to_matching_cell(importer, 'Product description', f"{data['species']} - {data['code']} - {data['weights'][2]} kg")
    transport_cells = unique_cells(transport.rows[1])
    transport_values = [f"{data['transport'][0]}\nVessel / bill of lading {data['transport'][1]}", data['exporter'], f"Container {data['transport'][2]}", 'United Kingdom', IMPORTER]
    for cell, value in zip(transport_cells, transport_values):
        append_value(cell, value)
    document.save(output)


def set_following_paragraph(document, heading, value):
    for index, paragraph in enumerate(document.paragraphs[:-1]):
        if heading in paragraph.text:
            target = document.paragraphs[index + 1]
            target.text = value
            for run in target.runs:
                run.bold = True
                run.font.size = Pt(9)
            return
    raise ValueError(f'Could not find template paragraph: {heading}')


def build_processing_statement(data, output):
    document = Document(TEMPLATES / 'Processing Statement Sample.docx')
    for paragraph in document.paragraphs:
        if 'DOCUMENT NUMBER' in paragraph.text:
            paragraph.add_run(' ' + data['number']).bold = True
        elif 'I confirm that the processed fishery products:' in paragraph.text:
            paragraph.add_run(' ' + data['description']).bold = True
    set_following_paragraph(document, 'Name and address of the processing plant:', 'EUROPEAN SEAFOOD INVESTMENTS PORTUGAL S.A.\nPorto de Pesca, 2520-630 Peniche, Portugal')
    set_following_paragraph(document, 'Name and address of the exporter', 'European Seafood Exports Portugal S.A.\nPorto de Pesca, 2520-630 Peniche, Portugal')
    set_following_paragraph(document, 'Approval number of the processing plant:', 'PT C 1234 CE')
    set_following_paragraph(document, 'Health certificate number and date:', f"HC-PT-2026-1149 - 10 December 2026")
    for row, certificate in zip(document.tables[0].rows[1:], data['certificates']):
        for cell, value in zip(unique_cells(row), certificate):
            append_value(cell, value)
    responsibility = unique_cells(document.tables[1].rows[0])
    for cell, value in zip(responsibility, ['Quality Manager', 'Signed electronically', '10 December 2026', 'Peniche, Portugal']):
        append_value(cell, value)
    authority = unique_cells(document.tables[2].rows[0])
    for cell, value in zip(authority, ['Portuguese Directorate-General for Natural Resources, Safety and Maritime Services', 'Validated electronically', '10 December 2026', 'Lisbon, Portugal']):
        append_value(cell, value)
    document.save(output)


def build_bill_of_lading(data, output):
    document = canvas.Canvas(str(output), pagesize=(595, 842))
    document.setTitle('Sample bill of lading')
    document.setFont('Helvetica-Bold', 18)
    document.drawString(50, 790, 'BILL OF LADING')
    document.setFont('Helvetica', 9)
    document.drawRightString(545, 794, 'Fictional document for user research only')

    rows = [
        ('Document number', data['number']),
        ('Date issued', data['issued']),
        ('Shipper', data['shipper']),
        ('Consignee', data['consignee']),
        ('Port of loading', data['departure']),
        ('Port of discharge', data['destination']),
        ('Container number', data['container']),
        ('Description of goods', data['cargo']),
        ('Gross weight', data['weight'])
    ]
    y_position = 735
    for label, value in rows:
        document.setFont('Helvetica-Bold', 10)
        document.drawString(55, y_position, label)
        document.setFont('Helvetica', 10)
        document.drawString(190, y_position, value)
        document.line(50, y_position - 8, 545, y_position - 8)
        y_position -= 48

    document.setFont('Helvetica', 9)
    document.drawString(55, 270, 'Received for shipment in apparent good order and condition.')
    document.drawString(55, 245, 'Carrier signature: Signed electronically for sample')
    document.save()


def export_pdf(docx_path, pdf_path):
    script = '''
on run argv
  tell application "Microsoft Word"
        with timeout of 300 seconds
            open file name (item 1 of argv)
            set generatedDocument to active document
            set generatedDocumentName to name of generatedDocument
            save as generatedDocument file name (item 2 of argv) file format format PDF
            repeat with openDocument in (get every document)
                try
                    if name of openDocument is generatedDocumentName then close openDocument saving no
                end try
            end repeat
        end timeout
  end tell
end run
'''
    subprocess.run(['osascript', '-e', script, str(docx_path), str(pdf_path)], check=True)


def watermark(source, destination):
    reader = PdfReader(source)
    writer = PdfWriter()
    for page in reader.pages:
        width = float(page.mediabox.width)
        height = float(page.mediabox.height)
        overlay_buffer = io.BytesIO()
        overlay = canvas.Canvas(overlay_buffer, pagesize=(width, height))
        overlay.saveState()
        overlay.setFillColorRGB(0.82, 0.82, 0.82)
        overlay.setFont('Helvetica-Bold', 68)
        overlay.translate(width / 2, height / 2)
        overlay.rotate(45)
        overlay.drawCentredString(0, 0, 'SAMPLE')
        overlay.restoreState()
        overlay.save()
        overlay_buffer.seek(0)
        page.merge_page(PdfReader(overlay_buffer).pages[0])
        writer.add_page(page)
    with destination.open('wb') as pdf_file:
        writer.write(pdf_file)


def generate(pilot=False, bill_of_lading_only=False):
    records = [
        record for record in CATCH_CERTIFICATES + PROCESSING_STATEMENTS + BILLS_OF_LADING
        if record['filename'] in GENERATED_FILENAMES
    ]
    if bill_of_lading_only:
        records = BILLS_OF_LADING
    if pilot:
        records = records[:1]
    staging = VERSION_ROOT / '.tmp' / 'sample-evidence'
    if staging.exists():
        shutil.rmtree(staging)
    staging.mkdir(parents=True)
    with tempfile.TemporaryDirectory() as temporary:
        temporary_path = Path(temporary)
        for index, record in enumerate(records):
            docx_path = temporary_path / f"generated-{index}-{Path(record['filename']).stem}.docx"
            raw_pdf = temporary_path / record['filename']
            output_pdf = staging / record['filename']
            if record in CATCH_CERTIFICATES:
                build_catch_certificate(record, docx_path)
                export_pdf(docx_path, raw_pdf)
            elif record in PROCESSING_STATEMENTS:
                build_processing_statement(record, docx_path)
                export_pdf(docx_path, raw_pdf)
            else:
                build_bill_of_lading(record, raw_pdf)
            watermark(raw_pdf, output_pdf)
            print(output_pdf)
    return staging


def install(staging):
    for pdf in staging.glob('*.pdf'):
        shutil.copy2(pdf, TARGET / pdf.name)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--pilot', action='store_true', help='Generate only the first catch certificate')
    parser.add_argument('--bill-of-lading-only', action='store_true', help='Generate only the sample bill of lading')
    parser.add_argument('--install', action='store_true', help='Replace this version\'s evidence PDFs')
    arguments = parser.parse_args()
    generated = generate(arguments.pilot, arguments.bill_of_lading_only)
    if arguments.install:
        install(generated)